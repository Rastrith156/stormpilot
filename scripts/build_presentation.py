from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREENSHOT = ROOT / 'docs' / 'stormpilot-board.png'
REPORT_PATH = ROOT / 'docs' / 'StormPilot_Project_Report.docx'
PRESENTATION_PATH = ROOT / 'StormPilot_Presentation.pptx'
REPORT_DATE = '11 October 2026'
PRESENTATION_DATE = '12 October 2026'

NAVY = RGBColor(16, 24, 39)
PANEL = RGBColor(27, 37, 57)
PANEL_LIGHT = RGBColor(37, 50, 75)
WHITE = RGBColor(241, 245, 249)
MUTED = RGBColor(173, 187, 207)
CYAN = RGBColor(73, 214, 201)
BLUE = RGBColor(72, 160, 250)
AMBER = RGBColor(255, 200, 87)
RED = RGBColor(255, 91, 112)


def make_report():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = DocInches(0.65)
    section.bottom_margin = DocInches(0.65)
    section.left_margin = DocInches(0.75)
    section.right_margin = DocInches(0.75)

    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = DocPt(10)
    normal.font.color.rgb = DocRGB(42, 52, 66)
    for style_name, size in [('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[style_name]
        style.font.name = 'Aptos Display'
        style.font.size = DocPt(size)
        style.font.bold = True
        style.font.color.rgb = DocRGB(24, 52, 81)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run('StormPilot')
    run.bold = True
    run.font.name = 'Aptos Display'
    run.font.size = DocPt(30)
    run.font.color.rgb = DocRGB(24, 52, 81)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = DocPt(3)
    run = subtitle.add_run('Resilient Dispatch and Agile Delivery Lab')
    run.bold = True
    run.font.size = DocPt(16)
    meta = doc.add_paragraph(f'Project report | Prepared {REPORT_DATE} | Presentation {PRESENTATION_DATE}')
    meta.paragraph_format.space_after = DocPt(12)

    doc.add_heading('Executive Summary', level=1)
    doc.add_paragraph(
        'StormPilot is an educational local lab combining a simulated drone-dispatch API '
        'with a Scrum and Kanban project workflow. It demonstrates application health and '
        'metrics, SQLite-backed work tracking, container hardening, monitoring configuration, '
        'Kubernetes recovery patterns, and controlled fault-injection procedures. The project '
        'is implemented and has passed its local Python tests and smoke checks; the Docker, '
        'monitoring, and Kubernetes demonstrations still require validation in an environment '
        'with those tools available.'
    )

    doc.add_heading('Purpose and Scope', level=1)
    doc.add_paragraph(
        'The scenario models a dispatch service that should remain observable and recoverable '
        'during routine failures. Alongside the API, the web interface supports backlog stories, '
        'sprints, task assignment, priorities, due dates, and a four-column Kanban board with '
        'work-in-progress limits.'
    )

    doc.add_heading('System Design', level=1)
    doc.add_paragraph(
        'The browser UI calls a dependency-free Python HTTP server. API routes manage stories, '
        'sprints, tasks, board state, and WIP limits in SQLite. Operational routes expose health, '
        'readiness, dispatch, and Prometheus-formatted metrics. Docker Compose defines the API, '
        'Prometheus, and Grafana. Kubernetes manifests define a replicated API deployment, '
        'health probes, resource settings, a disruption budget, and an optional CPU HPA.'
    )
    for item in [
        'Application: Python standard-library HTTP server and JSON API; persistent SQLite data.',
        'Container: non-root runtime and HTTP health check; Compose API configured read-only with dropped capabilities.',
        'Observability: request counters and uptime metrics, Prometheus scrape/rules, provisioned Grafana dashboard.',
        'Resilience lab: Kubernetes reconciliation, guarded one-pod experiment, and optional metrics-server-dependent autoscaling.',
    ]:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('Local Demonstration', level=1)
    doc.add_paragraph(
        'Start the application from the repository root with `python app/server.py`, then open '
        'http://localhost:8080/ui. The captured local board shows the existing task “Kanban” in '
        'To Do with Critical priority. No sample records were added for this report.'
    )
    if SCREENSHOT.exists():
        image_paragraph = doc.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(str(SCREENSHOT), width=DocInches(6.4))
        caption = doc.add_paragraph('Figure 1. StormPilot Kanban board captured from the running local application.')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = DocPt(9)

    doc.add_heading('Validation Evidence', level=1)
    evidence = [
        ('Python integration tests', 'All six HTTP integration tests passed.'),
        ('Live API smoke checks', 'Identity, readiness, health, dispatch, and metrics routes responded.'),
        ('Bounded local traffic check', '130 successful requests and zero failures during a two-second smoke run; this is not a capacity or uptime claim.'),
        ('Configuration checks', 'Bash syntax and all YAML and JSON configuration files parsed successfully.'),
        ('Browser check', 'Kanban view displays the existing task; Scrum/Kanban switching and modal hidden-state behavior verified locally.'),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Shading Accent 1'
    table.rows[0].cells[0].text = 'Check'
    table.rows[0].cells[1].text = 'Recorded result'
    for check, result in evidence:
        cells = table.add_row().cells
        cells[0].text = check
        cells[1].text = result

    doc.add_heading('Limitations and Next Steps', level=1)
    doc.add_paragraph(
        'Docker/Compose builds, runtime monitoring and Grafana visual checks, Prometheus rule '
        'validation, Kubernetes server-side validation and pod recovery, Trivy scanning, SBOM '
        'generation, and registry publishing were not executed in the recorded environment. '
        'GitHub repository creation was rejected with HTTP 403, so no repository or image was '
        'published. Actual HPA behavior requires metrics-server and has not been tested. This '
        'project is an educational lab, not a production service or an availability certification.'
    )
    for item in [
        'Run Compose and verify API, Prometheus targets, and the Grafana dashboard.',
        'Use a disposable kind cluster to verify readiness and the guarded single-pod recovery exercise.',
        'Capture actual CI, scan, replacement-pod, and HPA evidence before making release or reliability claims.',
    ]:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('Conclusion', level=1)
    doc.add_paragraph(
        'StormPilot provides a coherent local demonstration of product workflow and reliability '
        'engineering practices. The strongest presentation is an evidence-led walkthrough: show '
        'the working board and local checks, then identify container and Kubernetes experiments '
        'as configured next steps until they have been run.'
    )
    doc.add_paragraph('Evidence basis: repository README, VALIDATION.md, docs/DEMO.md, docs/RUNBOOK.md, Dockerfile, and the local browser capture.')
    doc.save(REPORT_PATH)


def add_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def add_text(slide, text, x, y, w, h, size=18, color=WHITE, bold=False,
             font='Aptos', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_box(slide, x, y, w, h, fill=PANEL, line=PANEL_LIGHT, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_header(slide, title, kicker):
    add_text(slide, kicker.upper(), 0.65, 0.35, 11.9, 0.25, 10, CYAN, True)
    add_text(slide, title, 0.65, 0.72, 12.0, 0.55, 27, WHITE, True, 'Aptos Display')
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.38), Inches(0.72), Inches(0.045))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN
    bar.line.fill.background()


def add_bullet_list(slide, items, x, y, w, line_h=0.56, size=16, color=WHITE, bullet_color=CYAN):
    for index, item in enumerate(items):
        yy = y + index * line_h
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.11), Inches(0.11), Inches(0.11))
        dot.fill.solid()
        dot.fill.fore_color.rgb = bullet_color
        dot.line.fill.background()
        add_text(slide, item, x + 0.25, yy, w - 0.25, line_h - 0.02, size, color)


def new_slide(prs, title=None, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    if title:
        add_header(slide, title, kicker or 'StormPilot')
    return slide


def make_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = new_slide(prs)
    add_text(slide, 'STORMPILOT / PROJECT BRIEF', 0.8, 0.72, 8.8, 0.3, 11, CYAN, True)
    add_text(slide, 'Resilient Dispatch\n& Agile Delivery', 0.8, 1.45, 8.5, 1.65, 34, WHITE, True, 'Aptos Display')
    add_text(slide, 'A local reliability-engineering lab with a working Scrum and Kanban workflow',
             0.82, 3.38, 7.6, 0.7, 18, MUTED)
    add_text(slide, f'Project presentation  |  {PRESENTATION_DATE}', 0.82, 6.65, 8.0, 0.3, 12, MUTED)
    for x, number, label, color in [
        (9.55, '01', 'API + UI', CYAN),
        (10.62, '02', 'OBSERVE', BLUE),
        (11.70, '03', 'RECOVER', AMBER),
    ]:
        add_box(slide, x, 1.7, 0.95, 3.8, PANEL, PANEL_LIGHT)
        add_text(slide, number, x + 0.12, 2.0, 0.72, 0.45, 22, color, True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.1, 4.72, 0.76, 0.45, 10, WHITE, True, align=PP_ALIGN.CENTER)

    slide = new_slide(prs, 'Product workflow in the browser', '01 / User experience')
    if SCREENSHOT.exists():
        slide.shapes.add_picture(str(SCREENSHOT), Inches(0.65), Inches(1.65), width=Inches(8.15), height=Inches(4.15))
    add_box(slide, 9.1, 1.65, 3.55, 4.15)
    add_text(slide, 'LIVE LOCAL CAPTURE', 9.42, 1.98, 2.9, 0.25, 10, CYAN, True)
    add_text(slide, 'Kanban board', 9.42, 2.48, 2.9, 0.45, 21, WHITE, True, 'Aptos Display')
    add_bullet_list(slide, [
        'Four workflow stages with WIP limits',
        'Existing task: Kanban / To Do / Critical',
        'Scrum backlog and sprint views also available',
        'SQLite-backed tasks, stories, and sprints',
    ], 9.42, 3.18, 2.95, 0.58, 14)
    add_text(slide, 'Captured from localhost; no demo records added for this deck.', 0.68, 6.18, 8.8, 0.3, 10, MUTED)

    slide = new_slide(prs, 'One application, several reliability layers', '02 / Architecture')
    nodes = [
        (0.65, 2.0, 2.35, 1.15, 'Browser UI', 'Scrum + Kanban', CYAN),
        (3.55, 2.0, 2.55, 1.15, 'Python HTTP API', 'JSON + health + metrics', BLUE),
        (6.7, 2.0, 2.3, 1.15, 'SQLite', 'Persistent work data', AMBER),
        (9.6, 2.0, 3.05, 1.15, 'Prometheus / Grafana', 'Configured observability', CYAN),
    ]
    for x, y, w, h, title, subtitle, accent in nodes:
        add_box(slide, x, y, w, h)
        add_text(slide, title, x + 0.15, y + 0.2, w - 0.3, 0.32, 16, accent, True)
        add_text(slide, subtitle, x + 0.15, y + 0.64, w - 0.3, 0.28, 11, MUTED)
    for x in (3.05, 6.1, 9.0):
        add_text(slide, '->', x, 2.35, 0.4, 0.35, 20, WHITE, True, align=PP_ALIGN.CENTER)
    add_box(slide, 1.05, 4.15, 5.15, 1.25, PANEL, PANEL_LIGHT)
    add_text(slide, 'Docker Compose', 1.32, 4.4, 2.0, 0.28, 15, BLUE, True)
    add_text(slide, 'Local API + Prometheus + Grafana demo', 1.32, 4.84, 4.5, 0.27, 12, MUTED)
    add_box(slide, 7.05, 4.15, 5.15, 1.25, PANEL, PANEL_LIGHT)
    add_text(slide, 'Kubernetes', 7.32, 4.4, 2.0, 0.28, 15, AMBER, True)
    add_text(slide, 'Replicas, probes, reconciliation, optional HPA', 7.32, 4.84, 4.55, 0.27, 12, MUTED)
    add_text(slide, 'Compose and Kubernetes layers are configured; end-to-end execution is pending in this environment.',
             0.75, 6.15, 11.8, 0.4, 11, MUTED)

    slide = new_slide(prs, 'Reliability controls are designed into the lab', '03 / Engineering controls')
    controls = [
        ('01', 'Container boundary', 'Non-root image; read-only API filesystem and dropped capabilities in Compose.', CYAN),
        ('02', 'Health and signals', 'Health/readiness endpoints, request counters, uptime metric, Prometheus rules.', BLUE),
        ('03', 'Recovery patterns', 'Kubernetes replicas, probes, disruption budget, and guarded one-pod experiment.', AMBER),
        ('04', 'Scaling path', 'CPU HPA manifest is optional and depends on a working metrics-server.', RED),
    ]
    for index, (number, title, description, accent) in enumerate(controls):
        y = 1.7 + index * 1.12
        add_box(slide, 0.72, y, 11.88, 0.88)
        add_text(slide, number, 0.98, y + 0.2, 0.55, 0.28, 16, accent, True)
        add_text(slide, title, 1.72, y + 0.14, 2.55, 0.28, 15, WHITE, True)
        add_text(slide, description, 4.25, y + 0.14, 7.95, 0.5, 13, MUTED)
    add_text(slide, 'Configured controls are not the same as a completed failure-recovery test.',
             0.78, 6.45, 11.5, 0.3, 12, AMBER, True)

    slide = new_slide(prs, 'What has been verified locally', '04 / Test evidence')
    metrics = [
        (0.72, '6 / 6', 'HTTP integration tests passed', CYAN),
        (4.53, '130', 'successful requests in a bounded 2-second smoke run', BLUE),
        (8.34, '0', 'failures in that smoke run', AMBER),
    ]
    for x, value, label, accent in metrics:
        add_box(slide, x, 1.75, 3.48, 1.55)
        add_text(slide, value, x + 0.25, 1.98, 2.95, 0.54, 28, accent, True, 'Aptos Display')
        add_text(slide, label, x + 0.25, 2.63, 2.95, 0.45, 12, MUTED)
    add_bullet_list(slide, [
        'Live checks covered identity, readiness, health, dispatch, and metrics.',
        'Bash syntax and YAML/JSON configuration parsing passed.',
        'Browser confirms the task board and modal/tab navigation locally.',
    ], 1.0, 4.05, 11.2, 0.63, 15)
    add_text(slide, 'Short smoke test only: not a capacity, uptime, or SLO measurement.',
             1.0, 6.3, 11.2, 0.35, 12, AMBER, True)

    slide = new_slide(prs, 'Next steps before claiming full-stack resilience', '05 / Limits and demo plan')
    add_box(slide, 0.72, 1.7, 5.75, 3.95)
    add_text(slide, 'Not yet run', 1.02, 2.0, 4.9, 0.42, 20, RED, True, 'Aptos Display')
    add_bullet_list(slide, [
        'Docker/Compose build and monitoring runtime',
        'Grafana visual QA and Prometheus rule validation',
        'kind pod replacement, HPA, and image scan',
        'SBOM and registry publication',
    ], 1.02, 2.72, 5.0, 0.62, 14, MUTED, RED)
    add_box(slide, 6.82, 1.7, 5.75, 3.95)
    add_text(slide, 'Presentation flow', 7.12, 2.0, 4.9, 0.42, 20, CYAN, True, 'Aptos Display')
    add_bullet_list(slide, [
        'Show Scrum/Kanban board and the saved task',
        'Run the six tests and local API smoke checks',
        'Explain Compose observability and Kubernetes recovery design',
        'Separate executed evidence from configured experiments',
    ], 7.12, 2.72, 5.0, 0.62, 14, MUTED)
    add_text(slide, 'GitHub repository creation was blocked by HTTP 403; no remote publication occurred.',
             0.82, 6.22, 11.7, 0.42, 12, AMBER, True)
    add_text(slide, 'StormPilot is an educational lab, not a production deployment.',
             0.82, 6.68, 11.7, 0.32, 11, MUTED)

    prs.save(PRESENTATION_PATH)


if __name__ == '__main__':
    if not SCREENSHOT.exists():
        raise SystemExit(f'Missing screenshot: {SCREENSHOT}')
    make_report()
    make_presentation()
    print(f'Created {REPORT_PATH}')
    print(f'Created {PRESENTATION_PATH}')
